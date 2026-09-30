"""Generate the Agri AI Network pitch deck (12 slides, 16:9).

    uv run --with python-pptx docs/pitch/build_deck.py

Writes docs/pitch/agri-ai-network-pitch.pptx. Screenshots come from docs/pitch/img/ (captured from the running
apps). Structured around the hackathon evaluation criteria (PRD §55) plus cross-border / BRICS portability.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "agri-ai-network-pitch.pptx"
IMG = HERE / "img"

GREEN = RGBColor(0x15, 0x80, 0x3D)
DARK = RGBColor(0x14, 0x1F, 0x17)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x65, 0x70)
LIGHT = RGBColor(0xF1, 0xF7, 0xF2)
AMBER = RGBColor(0xB4, 0x53, 0x09)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

URL_FARMER = "agri-ai-network-farmer-web-847963771142.asia-south1.run.app"
URL_OFFICER = "agri-ai-network-officer-dashboard-847963771142.asia-south1.run.app"
URL_REPO = "github.com/ipavanreddy/agri-ai-network"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, runs, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         spacing=1.1):
    """runs: str | list[str | tuple[str, dict]] - one paragraph per list item."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    for i, item in enumerate([runs] if isinstance(runs, str) else runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        content, opts = (item, {}) if isinstance(item, str) else item
        r = p.add_run()
        r.text = content
        f = r.font
        f.name = FONT
        f.size = Pt(opts.get("size", size))
        f.bold = opts.get("bold", bold)
        f.color.rgb = opts.get("color", color)
        if opts.get("space_after") is not None:
            p.space_after = Pt(opts["space_after"])
    return box


def bullets(slide, x, y, w, h, items, size=17, color=INK, gap=8):
    return text(slide, x, y, w, h, [(f"•  {t}", {"space_after": gap}) for t in items], size=size, color=color)


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def header(slide, title, kicker=None, weight=None):
    rect(slide, 0, 0, W, Inches(0.12), GREEN)
    if kicker:
        text(slide, Inches(0.6), Inches(0.35), Inches(9), Inches(0.4), kicker.upper(), size=12, color=GREEN, bold=True)
    text(slide, Inches(0.6), Inches(0.65), Inches(10.4), Inches(0.9), title, size=30, color=DARK, bold=True)
    if weight:
        pill = rect(slide, Inches(11.2), Inches(0.5), Inches(1.6), Inches(0.5), GREEN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        pill.text_frame.text = weight
        para = pill.text_frame.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        para.runs[0].font.size, para.runs[0].font.bold = Pt(14), True
        para.runs[0].font.color.rgb, para.runs[0].font.name = WHITE, FONT


def footer(slide, n):
    text(slide, Inches(0.6), Inches(7.05), Inches(8), Inches(0.3), "Agri AI Network · Build with AI · Track 4", size=10,
         color=MUTED)
    text(slide, Inches(11.9), Inches(7.05), Inches(0.9), Inches(0.3), str(n), size=10, color=MUTED, align=PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, title, body, accent=GREEN, size=14):
    rect(slide, x, y, w, h, LIGHT)
    rect(slide, x, y, Inches(0.08), h, accent)
    text(slide, x + Inches(0.25), y + Inches(0.15), w - Inches(0.4), Inches(0.45), title, size=17, bold=True, color=DARK)
    body_items = body if isinstance(body, list) else [body]
    text(slide, x + Inches(0.25), y + Inches(0.62), w - Inches(0.4), h - Inches(0.7),
         [(b, {"space_after": 4}) for b in body_items], size=size, color=INK)


def picture(slide, path, x, y, max_w, max_h):
    if not path.exists():
        return
    pic = slide.shapes.add_picture(str(path), x, y)
    scale = min(max_w / pic.width, max_h / pic.height)
    pic.width, pic.height = int(pic.width * scale), int(pic.height * scale)
    pic.left = int(x + (max_w - pic.width) / 2)
    pic.line.color.rgb = RGBColor(0xD1, 0xD5, 0xDB)


def notes(slide, t):
    slide.notes_slide.notes_text_frame.text = t


n = 0


def new_slide():
    global n
    n += 1
    s = prs.slides.add_slide(BLANK)
    if n > 1:
        footer(s, n)
    return s


# 1 - Title ------------------------------------------------------------------------------------------------------
s = new_slide()
rect(s, 0, 0, W, H, DARK)
rect(s, 0, Inches(5.2), W, Inches(0.08), GREEN)
text(s, Inches(0.8), Inches(1.3), Inches(11.5), Inches(0.5), "BUILD WITH AI · TRACK 4", size=14, color=RGBColor(0x86, 0xEF, 0xAC), bold=True)
text(s, Inches(0.8), Inches(1.9), Inches(11.5), Inches(1.2), "Agri AI Network", size=58, color=WHITE, bold=True)
text(s, Inches(0.8), Inches(3.1), Inches(11.5), Inches(1.6),
     ["One AI advisor for every smallholder's own field,",
      "one interoperable agriculture data network for every state."], size=26, color=RGBColor(0xD1, 0xFA, 0xE5))
text(s, Inches(0.8), Inches(5.5), Inches(11.8), Inches(1.2),
     [f"Farmer app  ·  {URL_FARMER}", f"Officer dashboard  ·  {URL_OFFICER}", f"Code  ·  {URL_REPO}"],
     size=14, color=RGBColor(0xE5, 0xE7, 0xEB))
notes(s, "Open with the one-line promise. Live URLs on screen; the demo follows the PRD §44 story.")

# 2 - Problem -----------------------------------------------------------------------------------------------------
s = new_slide()
header(s, "Farm decisions ignore data that already exists", "Problem", "Fit 20%")
bullets(s, Inches(0.6), Inches(1.7), Inches(6.4), Inches(4.8), [
    "86 % of India's ~146 million farm holdings are small or marginal (Agriculture Census 2015-16).",
    "Weather forecasts, Soil Health Cards, satellite imagery and state records exist, but in separate systems, "
    "formats and languages.",
    "Advice arrives generic, late, in English, or not at all; a leaf spot is diagnosed by guesswork.",
    "Every state builds its own portal, so nothing is shared and officers see districts, not fields.",
], size=18)
card(s, Inches(7.4), Inches(1.8), Inches(5.3), Inches(1.45), "The farmer's questions",
     "Is my crop healthy? Should I irrigate? What is this spot on the leaf? What should I sow next?")
card(s, Inches(7.4), Inches(3.45), Inches(5.3), Inches(1.45), "The officer's questions",
     "Which districts are at risk this week? Where are disease hotspots? Can I trust the AI advice going out?",
     accent=AMBER)
card(s, Inches(7.4), Inches(5.1), Inches(5.3), Inches(1.45), "The state's question",
     "Can we join a national network without rebuilding our systems?", accent=MUTED)
notes(s, "Problem-solution fit (20%): localized intelligence, crop risk, weather uncertainty, disease, climate resilience.")

# 3 - Solution ----------------------------------------------------------------------------------------------------
s = new_slide()
header(s, "Farm → Data → AI → Advisory → Diagnosis → Action", "Solution", "Fit 20%")
steps = [("1  Field", "Search the village, tap the field corners, pick the crop"),
         ("2  Intelligence", "Live weather, Soil Health Card, SoilGrids, Sentinel-2 NDVI"),
         ("3  Farm Health", "Explainable 0-100 score; missing data is never invented"),
         ("4  Advisory", "Gemini: observed → interpretation → action + regenerative practice"),
         ("5  Crop Doctor", "Leaf photo → Gemini multimodal → condition, confidence, nearest KVK"),
         ("6  Voice", "Ask in Telugu or Hindi, hear the answer grounded in the field")]
cw, ch = Inches(3.95), Inches(1.55)
for i, (t, b) in enumerate(steps):
    card(s, Inches(0.6) + (cw + Inches(0.2)) * (i % 3), Inches(1.8) + (ch + Inches(0.25)) * (i // 3), cw, ch, t, b, size=15)
text(s, Inches(0.6), Inches(5.5), Inches(12.1), Inches(1.2),
     ["Two roles only: the Farmer acts on advice; the Agriculture Officer sees the region and approves every AI output "
      "before it counts."], size=17, color=MUTED)
notes(s, "This is the single continuous journey shown in the demo video.")

# 4 - Product: farmer ---------------------------------------------------------------------------------------------
s = new_slide()
header(s, "The farmer sees one field, one score, one next action", "Product · farmer app")
bullets(s, Inches(0.6), Inches(1.7), Inches(5.6), Inches(5), [
    "Every card shows its source and freshness, with Live or Sample · synthetic badges.",
    "Farm Health = Vegetation 30 % + Soil 25 % + Weather 20 % + Water 15 % + Crop condition 10 %.",
    "Advisory in English, हिन्दी, తెలుగు; the same advice is translated, not re-analysed.",
    "Crop Doctor links straight to the nearest agriculture office with road distance and drive time.",
    "Works offline for demos: 'Use sample scenario data' replays a labelled dry-spell story.",
], size=17)
picture(s, IMG / "farmer.png", Inches(6.5), Inches(1.45), Inches(6.3), Inches(5.5))
notes(s, "Screenshot from the running app (OpenStreetMap basemap in the screenshot; Google Maps satellite in production).")

# 5 - AI and technical execution ----------------------------------------------------------------------------------
s = new_slide()
header(s, "Gemini reasons over facts; it never invents them", "AI / technical execution", "AI 25%")
card(s, Inches(0.6), Inches(1.7), Inches(4.0), Inches(2.4), "Structured orchestration", [
    "Canonical field context in, Pydantic-validated JSON out (generate_structured).",
    "Every record stores model_name, model_version, prompt_version (ai/prompts/*_v1.md)."], size=14)
card(s, Inches(4.75), Inches(1.7), Inches(4.0), Inches(2.4), "Gemini 2.5 Flash on Vertex AI", [
    "Advisory, crop-option explanation, voice Q&A.",
    "Multimodal Crop Doctor: condition, symptoms, likelihoods, confidence, expert-review flag."], size=14)
card(s, Inches(8.9), Inches(1.7), Inches(3.85), Inches(2.4), "Google AI services", [
    "Speech-to-Text + Text-to-Speech (en/hi/te-IN).",
    "Cloud Translation for localised advisories.",
    "Earth Engine Sentinel-2 NDVI/NDMI per field polygon."], size=14)
card(s, Inches(0.6), Inches(4.35), Inches(6.0), Inches(2.3), "Grounding and safety", [
    "Observed data → AI interpretation → recommendation kept as separate fields.",
    "Deterministic Farm Health rules; Gemini explains, rules score.",
    "Crop Doctor is always preliminary, with a disclaimer and a human officer review."], accent=AMBER, size=14)
card(s, Inches(6.75), Inches(4.35), Inches(6.0), Inches(2.3), "Honest fallbacks", [
    "Each integration is probed with a real call; the UI shows what is live and what is on a labelled fallback.",
    "Without keys the full journey still runs on deterministic rules and labelled sample data."], accent=MUTED, size=14)
notes(s, "AI/technical execution (25%): Gemini, Gemini multimodal, Vertex AI, Earth Engine, structured orchestration.")

# 6 - Officer -----------------------------------------------------------------------------------------------------
s = new_slide()
header(s, "Officers see state → district → block, and approve the AI", "Product · officer dashboard")
picture(s, IMG / "officer.png", Inches(0.6), Inches(1.5), Inches(7.4), Inches(5.3))
bullets(s, Inches(8.3), Inches(1.7), Inches(4.5), Inches(5), [
    "Regional risk map, crop distribution, weather risk, water stress and disease / stress hotspots.",
    "Indicators served from BigQuery; platform activity live from the field records.",
    "Review queue: every advisory, crop plan and diagnosis starts as pending; the farmer sees approve or reject.",
    "Interoperability tab: raw state record → adapter → identical canonical record.",
], size=16)
notes(s, "FR-09 and the human-in-the-loop requirement for consequential actions.")

# 7 - Depth and reach across India ----------------------------------------------------------------------------------
s = new_slide()
header(s, "A new state is a config file and a data export, not a rebuild", "Depth & reach across India", "Reach 20%")
flow = ["State export", "Declarative adapter\n(data/adapters/*.json)", "Canonical schema\n(Pydantic → JSON Schema)",
        "Shared APIs + AI"]
bw = Inches(2.75)
for i, label in enumerate(flow):
    x = Inches(0.6) + (bw + Inches(0.38)) * i
    b = rect(s, x, Inches(1.75), bw, Inches(1.1), GREEN if i == 3 else LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, Inches(1.75), bw, Inches(1.1), label, size=15, bold=True, color=WHITE if i == 3 else DARK,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < 3:
        text(s, x + bw, Inches(1.95), Inches(0.38), Inches(0.6), "→", size=26, color=GREEN, align=PP_ALIGN.CENTER)
states = [("Andhra Pradesh", "Groundnut · e-Crop + Soil Health Card format · Telugu · codes GN, RG"),
          ("Maharashtra", "Cotton · MahaDBT format · Marathi → Hindi fallback · codes kapus, tur"),
          ("Punjab", "Wheat · land records + PAU soil lab · kanal → acres · codes kanak, sarson")]
for i, (t, b) in enumerate(states):
    card(s, Inches(0.6) + Inches(4.1) * i, Inches(3.2), Inches(3.95), Inches(1.5), t, b, size=14)
bullets(s, Inches(0.6), Inches(5.0), Inches(12.2), Inches(1.9), [
    "Adapters declare field paths, crop code tables, unit conversions (ha/kanal → acres, g/kg → %, kg/acre → kg/ha), "
    "date formats and language fallbacks. One engine applies them all.",
    "Farm Health, Gemini advisory, Crop Doctor, voice and the officer view run unchanged for every state.",
], size=16)
notes(s, "Depth & reach (20%): state-independent model, multi-state config, multi-language, reusable APIs.")

# 8 - Deployability and scalability --------------------------------------------------------------------------------
s = new_slide()
header(s, "Cloud-native: one command deploys the network", "Deployability & scalability", "Deploy 20%")
card(s, Inches(0.6), Inches(1.7), Inches(6.0), Inches(2.5), "Running on Google Cloud", [
    "3 Cloud Run services (API + 2 Next.js apps), asia-south1, built by Cloud Build.",
    "Secrets in Secret Manager; Vertex AI, BigQuery and Cloud Storage via the service account.",
    "infrastructure/cloud-run/deploy.sh: repeatable; Gemini key attached automatically when it exists."], size=14)
card(s, Inches(6.75), Inches(1.7), Inches(6.0), Inches(2.5), "API-first, modular", [
    "Documented REST API (/docs), canonical JSON Schemas, versioned prompts.",
    "Domain modules (farms, intelligence, advisory, crops, diagnosis, localization, interop, analytics) "
    "can split into separate services when load demands."], size=14)
phases = [("Hackathon", "3 states · 3 languages · real + labelled sample data"),
          ("District pilot", "1 state, many districts, government feeds"),
          ("State rollout", "State models, officer workflows, more languages"),
          ("National network", "Shared contracts, model registry, inter-state exchange")]
pw = Inches(2.95)
for i, (t, b) in enumerate(phases):
    x = Inches(0.6) + (pw + Inches(0.13)) * i
    rect(s, x, Inches(4.5), pw, Inches(0.5), GREEN if i == 0 else DARK)
    text(s, x, Inches(4.5), pw, Inches(0.5), t, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + Inches(0.1), Inches(5.1), pw - Inches(0.2), Inches(1.3), b, size=14, color=INK)
notes(s, "Deployability (20%): cloud-native services, API-first, adapters, standard schemas, modular AI, deployed app.")

# 9 - Impact ------------------------------------------------------------------------------------------------------
s = new_slide()
header(s, "Impact: better decisions on millions of small fields", "Impact potential", "Impact 15%")
impacts = [("Crop resilience", "Early NDVI dips and weather risk flagged before yield is lost."),
           ("Water efficiency", "Irrigate only when 14-day rain and forecast say so."),
           ("Disease response", "Photo to preliminary diagnosis in seconds, plus the nearest KVK."),
           ("Soil health", "Soil Health Card data turned into one regenerative practice per field."),
           ("Climate adaptation", "Crop options scored for local climate risk and rotation value."),
           ("Officer reach", "One dashboard across districts, with human approval of AI advice.")]
for i, (t, b) in enumerate(impacts):
    card(s, Inches(0.6) + (cw + Inches(0.2)) * (i % 3), Inches(1.7) + Inches(1.75) * (i // 3), cw, Inches(1.55), t, b,
         size=15)
text(s, Inches(0.6), Inches(5.35), Inches(12.2), Inches(1.4),
     ["What we will measure in a pilot: advisories acted on, irrigation events avoided, time from symptom to officer "
      "contact, share of AI outputs approved, farmers reached per officer."], size=16, color=MUTED)
notes(s, "Impact (15%): small and marginal farmers; resilience, water, disease, soil, climate.")

# 10 - Cross-border / BRICS -----------------------------------------------------------------------------------------
s = new_slide()
header(s, "Built for India, portable across borders", "Cross-border · BRICS portability")
bullets(s, Inches(0.6), Inches(1.7), Inches(6.2), Inches(5), [
    "Nothing in the core is India-specific: the canonical schema, Farm Health rules and AI prompts take any "
    "country's records through an adapter.",
    "Language is configuration: add Portuguese, Russian, Chinese, Arabic or Swahili strings and speech locales.",
    "Global public data already in the loop: Sentinel-2 (Earth Engine), Open-Meteo, ISRIC SoilGrids.",
    "Each partner keeps its data sovereign and runs its own Cloud Run deployment; only schemas and models are shared.",
], size=17)
card(s, Inches(7.2), Inches(1.8), Inches(5.5), Inches(1.5), "Same adapter pattern", "Brazil's state rural registries, "
     "South Africa's provincial extension data or Ethiopia's woreda soil maps map to the same canonical records.",
     size=14)
card(s, Inches(7.2), Inches(3.5), Inches(5.5), Inches(1.5), "Shared intelligence, local control",
     "Common data contracts enable cross-border pest and weather early warnings without moving farmer data.",
     accent=AMBER, size=14)
card(s, Inches(7.2), Inches(5.2), Inches(5.5), Inches(1.4), "Global South scale",
     "Smallholders produce a large share of food in BRICS+ countries; one open architecture serves them all.",
     accent=MUTED, size=14)
notes(s, "Portability story: adapters + canonical schema + config-driven language.")

# 11 - Status: live vs next -----------------------------------------------------------------------------------------
s = new_slide()
header(s, "What runs today, and what switches on next", "Status and roadmap")
card(s, Inches(0.6), Inches(1.7), Inches(6.0), Inches(3.9), "Live on real Google services", [
    "Gemini 2.5 Flash on Vertex AI (advisory, crop options, Crop Doctor, voice Q&A)",
    "Cloud Speech-to-Text, Text-to-Speech, Translation",
    "Maps Geocoding, Places, Routes + Maps JavaScript basemap",
    "BigQuery regional indicators · Cloud Storage crop photos",
    "Open-Meteo weather, ISRIC SoilGrids texture",
    "Cloud Run, Cloud Build, Artifact Registry, Secret Manager"], size=15)
card(s, Inches(6.75), Inches(1.7), Inches(6.0), Inches(3.9), "Code complete, switching on next", [
    "Earth Engine Sentinel-2 NDVI: waiting for project registration (sample series until then, labelled).",
    "Firestore + Firebase Auth: waiting for Firebase on the project (SQLite fallback).",
    "Real state feeds replace the labelled sample exports: a data change, not a code change.",
    "Next: Marathi and Punjabi, crop-specific disease models on Vertex AI, SMS / IVR channel for feature phones."],
     accent=AMBER, size=15)
notes(s, "Be explicit about what is live vs demo; the apps show the same split in their header badges.")

# 12 - Close ------------------------------------------------------------------------------------------------------
s = new_slide()
rect(s, 0, 0, W, H, DARK)
text(s, Inches(0.8), Inches(1.2), Inches(11.8), Inches(1.0), "One field → one district → one state → a national network",
     size=30, color=WHITE, bold=True)
text(s, Inches(0.8), Inches(2.7), Inches(11.5), Inches(1.4),
     ["Agri AI Network turns India's scattered agriculture data into advice a farmer can act on today,",
      "and gives every state a way to join without rebuilding."], size=20, color=RGBColor(0xD1, 0xFA, 0xE5))
text(s, Inches(0.8), Inches(4.4), Inches(11.8), Inches(1.6),
     [f"Try it  ·  {URL_FARMER}", f"Officer view  ·  {URL_OFFICER}", f"Source  ·  {URL_REPO}"],
     size=16, color=RGBColor(0xE5, 0xE7, 0xEB))
text(s, Inches(0.8), Inches(6.4), Inches(11.8), Inches(0.5), "Thank you", size=16, color=RGBColor(0x86, 0xEF, 0xAC),
     bold=True)
footer(s, n)

prs.save(OUT)
print(f"wrote {OUT.relative_to(HERE.parents[1])} ({len(prs.slides)} slides, {Emu(W).inches:.2f}x{Emu(H).inches:.2f} in)")
